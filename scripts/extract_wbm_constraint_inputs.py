"""Extract the data inputs of the COBRA Toolbox whole-body constraint functions into one JSON file.

runIEM_HH.m (Harvey branch) calls, before any IEM:
    standardPhysiolDefaultParameters   (getOrganWeightFraction; blood-flow spreadsheet)
    physiologicalConstraintsHMDBbased  (HMDB blood, CSF and urine concentration tables)
    EUAverageDietNew; setDietConstraints  (diet table; AGORAEssentialMetabolites)

This script reads those inputs from a COBRA Toolbox checkout and writes them, with the toolbox
commit and per-file git blob hashes, to data/iem/wbm_constraint_inputs_v0.3.json. The constraint
logic itself is ported in gembench/wbm_constraints.py. Values are kept as MATLAB would read them:
  * HMDB tables: rows after the '####' header in file order (first match wins in MATLAB);
  * blood-flow percentages: the text cells hold quoted numbers such as "'0.05'", which the
    Toolbox strips with B(2:end-1) before str2num; an empty result means "use 1 percent";
  * organ weights: getOrganWeightFraction's column arithmetic on the saved xlsread output;
  * diet and AGORA lists: the literal cell arrays, after the Toolbox's regexprep renaming.

Usage: python scripts/extract_wbm_constraint_inputs.py [--toolbox external/cobratoolbox]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
from parse_iem_protocol import strip_matlab_comments  # noqa: E402

OUT = os.path.join(ROOT, "data", "iem", "wbm_constraint_inputs_v0.3.json")
WB = "src/analysis/wholeBody/PSCMToolbox"
FILES = {
    "blood_hmdb": f"{WB}/setConstraints/inputData/NormalBloodConcExtractedHMDB.txt",
    "csf_hmdb": f"{WB}/setConstraints/inputData/NormalCSFConcExtractedHMDB.txt",
    "urine_hmdb": f"{WB}/setConstraints/inputData/NormalUrineConcExtractedHMDB.txt",
    "blood_flow": f"{WB}/setConstraints/inputData/16_01_26_BloodFlowRatesPercentages.xlsx",
    "organ_weights": f"{WB}/setConstraints/inputData/Numbers_organWeightData_fromxls_16_01_29_OrganWeigths.mat",
    "diet": f"{WB}/setConstraints/diets/EUAverageDietNew.m",
    "agora_essential": f"{WB}/hostMicrobeInteraction/AGORAEssentialMetabolites.m",
    # code files whose logic is ported (hashes recorded so a toolbox change is visible)
    "physiological_constraints_code": f"{WB}/setConstraints/physiologicalConstraintsHMDBbased.m",
    "diet_constraints_code": f"{WB}/setConstraints/setDietConstraints.m",
    "default_parameters_code": f"{WB}/setConstraints/standardPhysiolDefaultParameters.m",
    "organ_weight_code": f"{WB}/setConstraints/organWeight/getOrganWeightFraction.m",
    "organ_lists_code": f"{WB}/io/OrganLists.m",
    "runiem_code": f"{WB}/runIEM_HH.m",
    "change_rxn_bounds_code": "src/analysis/FBA/changeRxnBounds.m",
}


def git_blob(path: str) -> str:
    data = open(path, "rb").read()
    return hashlib.sha1(b"blob %d\0" % len(data) + data).hexdigest()


def sha256(path: str) -> str:
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def read_hmdb(path: str) -> list[dict]:
    """Rows after the column header that follows the '####' line, in file order."""
    lines = open(path, newline="").read().replace("\r\n", "\n").split("\n")
    marks = [i for i, line in enumerate(lines) if "####" in line.split("\t")[0]]
    if len(marks) != 1:
        raise ValueError(f"{path}: expected one '####' header marker, found {len(marks)}")
    header = lines[marks[0] + 1].split("\t")
    vmh = [i for i, h in enumerate(header) if "VMH" in h]
    if len(vmh) != 1:
        raise ValueError(f"{path}: expected one VMH column")
    rows = []
    for line in lines[marks[0] + 2:]:
        if not line.strip():
            continue
        f = line.split("\t")
        if len(f) != len(header):
            raise ValueError(f"{path}: malformed row {line!r}")
        lo, hi = float(f[5]), float(f[6])
        if not (math.isfinite(lo) and math.isfinite(hi)):
            raise ValueError(f"{path}: non-finite concentration in {line!r}")
        if f[vmh[0]] != f[vmh[0]].strip() or not f[vmh[0]]:
            raise ValueError(f"{path}: VMH id with surrounding whitespace or empty in {line!r}")
        rows.append({"vmh": f[vmh[0]], "min": lo, "max": hi, "unit": f[4], "hmdb": f[0]})
    ids = [r["vmh"] for r in rows]
    if len(ids) != len(set(ids)):
        raise ValueError(f"{path}: duplicate VMH ids; MATLAB would use the first row of each")
    return rows


def matlab_quoted_number(text):
    """B = str2num(B(2:end-1)) on a text cell; None where MATLAB gets [] (use the 1 percent default)."""
    if text is None:
        return None
    inner = str(text)[1:-1].strip()
    try:
        return float(inner) if inner else None
    except ValueError:      # header text such as 'Blood flow percentage'; str2num gives [] there too
        return None


def read_blood_flow(path: str) -> dict:
    import openpyxl
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["BloodFlowPercentage"]
    grid = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]
    for r, row in enumerate(grid):
        if any(v == "Blood flow percentage" for v in row):
            perc_cols = [j for j, v in enumerate(row) if v == "Blood flow percentage"]
            organ_col = [j for j, v in enumerate(row) if v == "Organ"][0]
            break
    else:
        raise ValueError("no 'Blood flow percentage' header row")
    if any(not isinstance(grid[i][j], str) for i in range(len(grid)) for j in perc_cols if grid[i][j] is not None):
        raise ValueError("numeric percentage cells: xlsread's text output would be empty for them")
    organs = {}
    for row in grid:
        name = row[organ_col]
        if not isinstance(name, str):
            continue
        if name in organs:
            raise ValueError(f"duplicate organ row {name}")
        organs[name] = {"male_text": row[perc_cols[0]], "female_text": row[perc_cols[1]],
                        "male": matlab_quoted_number(row[perc_cols[0]]),
                        "female": matlab_quoted_number(row[perc_cols[1]])}
    return {"organ_column": organ_col + 1, "percentage_columns": [c + 1 for c in perc_cols], "organs": organs}


def read_organ_weights(path: str) -> dict:
    """getOrganWeightFraction.m on the saved xlsread output (no 'weigth' override)."""
    import numpy as np
    import scipy.io as sio
    d = sio.loadmat(path, squeeze_me=False)
    numbers, text = d["Numbers"], d["organWeightData"]

    def cell(i, j):
        v = text[i, j]
        if isinstance(v, np.ndarray):
            return "" if v.size == 0 else str(v.ravel()[0]) if v.dtype.kind in "US" else ""
        return str(v)

    cols = None
    organ_row = None
    for i in range(text.shape[0]):
        found = [j + 1 for j in range(text.shape[1]) if cell(i, j) == "% of body weight"]
        if found:
            cols = found
        if cell(i, 0) == "Body weight":
            organ_row = i + 1          # MATLAB organRow = i+1 (1-based) -> python index i+1
            break
    out = {}
    for sex, col in (("male", cols[0] - 2), ("female", cols[1] - 2)):
        names = [cell(i, 0) for i in range(organ_row, text.shape[0])]
        weights = numbers[1:, col - 2].tolist()
        fract = (numbers[1:, col - 1] / 100).tolist()
        if len(names) != len(weights):
            raise ValueError("organ names and weights differ in length")
        out[sex] = {"numbers_column": col, "body_weight_g": float(numbers[0, col - 2]),
                    "organs": [{"name": n, "weight_g": None if math.isnan(w) else w,
                                "fraction": None if math.isnan(f) else f} for n, w, f in zip(names, weights, fract)]}
    return out


def matlab_cell_strings(text: str, variable: str) -> list[str]:
    """Quoted strings of the cell array literal assigned to `variable`, comments removed."""
    body = strip_matlab_comments(text)
    m = re.search(rf"^\s*{variable}\s*=\s*\{{(.*?)\}}\s*;", body, flags=re.S | re.M)
    if not m:
        raise ValueError(f"cell array {variable} not found")
    return re.findall(r"'([^']*)'", m.group(1))


def read_diet(path: str) -> list[list[str]]:
    strings = matlab_cell_strings(open(path).read(), "Diet")
    if len(strings) % 2:
        raise ValueError("diet cell array does not have two columns")
    diet = []
    for rxn, value in zip(strings[0::2], strings[1::2]):
        float(value)   # str2num must succeed
        rxn = rxn.replace("EX_", "Diet_EX_").replace("(e)", "[d]")  # regexprep(Diet,'EX_','Diet_EX_'); '\(e\)' -> '[d]'
        diet.append([rxn, value])
    return diet


def read_agora(path: str) -> list[str]:
    names = matlab_cell_strings(open(path).read(), "AGORAessential")
    return [n.replace("EX_", "Diet_EX_").replace("[u]", "[d]") for n in names]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--toolbox", default=os.path.join(ROOT, "external", "cobratoolbox"))
    ap.add_argument("--out", default=OUT)
    args = ap.parse_args()
    tb = args.toolbox
    paths = {k: os.path.join(tb, v) for k, v in FILES.items()}
    commit = subprocess.run(["git", "-C", tb, "rev-parse", "HEAD"], capture_output=True, text=True).stdout.strip()
    provenance = {"toolbox_repository": "https://github.com/opencobra/cobratoolbox", "toolbox_commit": commit,
                  "files": {k: {"path": FILES[k], "git_blob_sha1": git_blob(p), "sha256": sha256(p)} for k, p in paths.items()}}
    for k, p in paths.items():
        head = subprocess.run(["git", "-C", tb, "rev-parse", f"HEAD:{FILES[k]}"], capture_output=True, text=True).stdout.strip()
        if head != provenance["files"][k]["git_blob_sha1"]:
            raise SystemExit(f"{FILES[k]} differs from the toolbox commit {commit}")
    data = {
        "schema": "wbm-constraint-inputs-v0.3",
        "provenance": provenance,
        "hmdb": {"blood": read_hmdb(paths["blood_hmdb"]), "csf": read_hmdb(paths["csf_hmdb"]),
                 "urine": read_hmdb(paths["urine_hmdb"])},
        "blood_flow": read_blood_flow(paths["blood_flow"]),
        "organ_weights": read_organ_weights(paths["organ_weights"]),
        "diet_eu_average": read_diet(paths["diet"]),
        "agora_essential": read_agora(paths["agora_essential"]),
    }
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with open(args.out, "w") as fh:
        json.dump(data, fh, indent=1, allow_nan=False)
        fh.write("\n")
    print(f"wrote {args.out}: blood {len(data['hmdb']['blood'])}, csf {len(data['hmdb']['csf'])}, urine "
          f"{len(data['hmdb']['urine'])} HMDB rows; {len(data['blood_flow']['organs'])} blood-flow rows; "
          f"diet {len(data['diet_eu_average'])}; AGORA essential {len(data['agora_essential'])}; toolbox {commit[:7]}")


if __name__ == "__main__":
    main()
